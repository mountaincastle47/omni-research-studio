"""Phase 6 Production Hardening & Operational Polish Suite Headless Verification.

Tests:
1. Output Token Continuation Loop & Markdown Auto-Repair (sanitize_unclosed_markdown)
2. Windows File Lock & Audio Stitching Defenses (safe_remove_file, silent buffer, mutagen ID3)
3. Google GenAI Explicit Context Caching (ContextCacheStore)
4. Universal Multi-Format Document Ingestion (.docx tables, CP949 chardet, Drag-and-Drop)
5. SQLite FTS5 Global Search Index & SearchDrawerWidget (BM25 ranking, <mark> snippet)
6. In-Page Find Bar (Ctrl+F), Hover Tooltips & BibTeX Bibliography Exporter
7. Source Authority Hierarchy Classification ([PRIMARY], [SECONDARY], [TERTIARY])
"""

from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock

from PyQt6.QtCore import QMimeData, QPointF, Qt, QUrl
from PyQt6.QtGui import QDragEnterEvent, QDropEvent
from PyQt6.QtWidgets import QApplication, QTextBrowser

from core.audio_studio import (
    SILENT_MP3_BUFFER_200MS,
    DialogueTurn,
    NeuralVoiceSynthesizer,
    PodcastScript,
    safe_remove_file,
)
from core.config import ConfigManager
from core.document_loader import DocumentCorpus, SourceDocument, load_docx
from core.exporter import build_bibtex_report, export_bibtex_report
from core.followup_engine import FollowUpEngine
from core.notebook_synthesis import (
    ContextCacheStore,
    NotebookSynthesizer,
    create_semantic_anchor_builder,
    sanitize_unclosed_markdown,
)
from core.search_index import CorpusSearchIndex
from core.storm_discovery import classify_source_authority
from ui.main_window import MainWindow
from ui.search_drawer import SearchDockWidget, SearchDrawerWidget
from ui.workspace_tabs import InPageFindBar

# Initialize headless Qt Application
os.environ["QT_QPA_PLATFORM"] = "offscreen"
app = QApplication.instance() or QApplication(sys.argv)


def test_1_markdown_repair_and_continuation():
    print("[1/7] Testing Output Token Continuation Loop & Markdown Auto-Repair...")

    # 1. Test unclosed code fence
    broken_code = "# Report\n```python\nprint('hello')\n"
    fixed_code = sanitize_unclosed_markdown(broken_code)
    assert fixed_code.count("```") % 2 == 0, "Unclosed code fence must be closed!"
    assert fixed_code.endswith("\n```\n"), "Closed fence should end with backticks"

    # 2. Test unclosed table row
    broken_table = "| Metric | Value |\n|---|---|\n| Fidelity | 99.9%"
    fixed_table = sanitize_unclosed_markdown(broken_table)
    assert fixed_table.strip().endswith("|"), "Unclosed table row must be closed with pipe!"

    # 3. Test continuation loop detection
    mock_client = MagicMock()
    mock_stream1 = [
        MagicMock(text="Part 1: Initial research telemetry... ", candidates=[MagicMock(finish_reason="FinishReason.MAX_TOKENS")]),
    ]
    mock_stream2 = [
        MagicMock(text="Part 2: Continuation and conclusion.", candidates=[MagicMock(finish_reason="FinishReason.STOP")]),
    ]
    mock_client.models.generate_content_stream.side_effect = [mock_stream1, mock_stream2]

    synth = NotebookSynthesizer(client=mock_client)
    doc = SourceDocument(cite_id=1, title="Test Source", content="Quantum neutral atom telemetry")
    res = synth.synthesize(query="Evaluate quantum telemetry", documents=[doc])

    assert "Part 1" in res.markdown_report
    assert "Part 2" in res.markdown_report
    assert mock_client.models.generate_content_stream.call_count == 2
    print("   -> Markdown auto-repair and output token continuation loop verified successfully!")


def test_2_audio_defenses_and_id3():
    print("[2/7] Testing Windows File Lock & Audio Stitching Defenses...")

    # 1. safe_remove_file
    with tempfile.NamedTemporaryFile(delete=False) as tf:
        tf.write(b"lock_test")
        tmp_path = Path(tf.name)
    assert safe_remove_file(tmp_path), "safe_remove_file should remove existing file"
    assert not tmp_path.exists(), "File should be deleted"

    # 2. 200ms silent MP3 frame buffer
    assert len(SILENT_MP3_BUFFER_200MS) == (417 * 8), "Silent MP3 buffer should contain 8 frames"
    assert SILENT_MP3_BUFFER_200MS.startswith(b"\xff\xfb\x90\x64"), "Buffer must start with valid MPEG Layer 3 sync header"

    # 3. ID3 tagging with mutagen
    with tempfile.TemporaryDirectory() as tmp_dir:
        synth = NeuralVoiceSynthesizer(language="en", output_dir=tmp_dir)
        script = PodcastScript(
            title="Quantum Systems Overview",
            model_used="test-model",
            turns=[
                DialogueTurn(speaker="Host A (Alex)", role_tag="Anchor", text="Welcome to OmniResearch."),
                DialogueTurn(speaker="Host B (Sam)", role_tag="Skeptic", text="Let us review the quantum findings."),
            ],
        )
        # Mock edge_tts
        dummy_chunk = b"\xff\xfb\x90\x64" + (b"\x00" * 413)
        synth._synthesize_async = AsyncMock(return_value=[dummy_chunk, dummy_chunk])

        out_path = synth.synthesize_script(script, force=True)
        assert out_path.exists()
        assert out_path.stat().st_size > len(dummy_chunk)

        from mutagen.id3 import ID3
        tags = ID3(str(out_path))
        assert tags.get("TIT2").text[0] == "Quantum Systems Overview"
        assert "Host A (Alex)" in tags.get("TPE1").text[0]
        assert tags.get("TALB").text[0] == "OmniResearch Studio"

    print("   -> Windows file lock defense, 200ms silence buffer, and mutagen ID3 tags verified successfully!")


def test_3_context_caching():
    print("[3/7] Testing Google GenAI Explicit Context Caching...")
    store = ContextCacheStore.get_instance()
    assert store is ContextCacheStore.get_instance(), "ContextCacheStore must be a singleton"

    h = "test_hash_123"
    model = "gemini-3.8-flash"
    store.put_cache(h, model, "cachedContents/omni_123", ttl_seconds=60.0)
    cached_name = store.get_cached_name(h, model)
    assert cached_name == "cachedContents/omni_123"

    # Expired check
    store.put_cache(h, model, "cachedContents/expired", ttl_seconds=-1.0)
    assert store.get_cached_name(h, model) is None, "Expired cache must return None"

    # Context caching trigger check in FollowUpEngine (>130,000 chars)
    mock_client = MagicMock()
    mock_cache_obj = MagicMock(name="cachedContents/followup_cache")
    mock_cache_obj.name = "cachedContents/followup_cache"
    mock_client.caches.create.return_value = mock_cache_obj
    mock_client.models.generate_content_stream.return_value = [MagicMock(text="Followup answer grounded in cache.")]

    engine = FollowUpEngine(client=mock_client)
    huge_corpus = [
        SourceDocument(cite_id=i, title=f"Huge Source {i}", content="Telemetry data " * 1500)
        for i in range(1, 12)
    ]
    ans = engine.stream_answer(question="What is the telemetry?", synthesis_report="Report", documents=huge_corpus)
    assert "Followup answer" in ans
    assert mock_client.caches.create.called, "Explicit context cache must be invoked for large context"

    print("   -> ContextCacheStore and explicit caching integration verified successfully!")


def test_4_universal_document_ingestion():
    print("[4/7] Testing Universal Multi-Format Document Ingestion (.docx, CP949, Drag-and-Drop)...")

    # 1. Test .docx ingestion with tables
    import docx
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tf:
        docx_path = Path(tf.name)

    doc = docx.Document()
    doc.add_paragraph("First technical paragraph.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Subsystem"
    table.cell(0, 1).text = "Status"
    table.cell(1, 0).text = "Neutral Atom Laser"
    table.cell(1, 1).text = "Locked"
    doc.save(str(docx_path))

    res = load_docx(docx_path)
    assert "<!--PAGE_START:1-->" in res["content"]
    assert "| Subsystem | Status |" in res["content"]
    assert "| Neutral Atom Laser | Locked |" in res["content"]

    corpus = DocumentCorpus()
    sdoc = corpus.ingest_file(docx_path)
    assert sdoc.source_type == "local_docx"
    assert "Subsystem" in sdoc.content
    docx_path.unlink()

    # 2. Test CP949 text detection via chardet
    with tempfile.NamedTemporaryFile(suffix=".txt", delete=False) as tf:
        txt_path = Path(tf.name)
    korean_raw = "양자 프로세서 제어 아키텍처 및 큐비트 충실도".encode("cp949")
    txt_path.write_bytes(korean_raw)

    sdoc2 = corpus.ingest_file(txt_path)
    assert sdoc2.source_type == "local_txt"
    assert "<!--PAGE_START:1-->" in sdoc2.content
    assert "양자 프로세서" in sdoc2.content
    txt_path.unlink()

    print("   -> Universal document ingestion (.docx tables, CP949 chardet) verified successfully!")


def test_5_sqlite_fts5_and_search_drawer():
    print("[5/7] Testing SQLite FTS5 Global Search Index & SearchDrawerWidget...")

    index = CorpusSearchIndex()
    doc = SourceDocument(
        cite_id=1,
        title="Neutral Atom Computing",
        content="<!--PAGE_START:1-->\n\nRydberg blockade mechanism creates Bell pairs.\n\n<!--PAGE_START:2-->\n\nCryogenic cooling reaches 4 Kelvin.",
    )
    index.index_source_document(doc)

    # BM25 Search
    matches = index.search_corpus("Rydberg")
    assert len(matches) == 1
    assert matches[0].doc_id == 1
    assert matches[0].page_num == 1
    assert "<mark>Rydberg</mark>" in matches[0].snippet

    matches_p2 = index.search_corpus("Cryogenic")
    assert len(matches_p2) == 1
    assert matches_p2[0].page_num == 2

    # UI Widget Test
    drawer = SearchDrawerWidget(search_index=index, lang="en")
    drawer.search_input.setText("Rydberg")
    drawer.perform_search()
    assert drawer.results_list.count() == 1

    selected_data = []
    drawer.result_selected.connect(lambda d, p: selected_data.append((d, p)))
    item = drawer.results_list.item(0)
    drawer._on_item_clicked(item)
    assert selected_data == [(1, 1)]

    dock = SearchDockWidget(search_index=index, lang="ko")
    assert dock.windowTitle() == "전문 검색 (FTS5)"

    index.close()
    print("   -> SQLite FTS5 search index, BM25 snippet ranking, and SearchDrawerWidget verified successfully!")


def test_6_find_bar_and_bibtex():
    print("[6/7] Testing In-Page Find Bar (Ctrl+F), Hover Tooltips & BibTeX Exporter...")

    # 1. InPageFindBar test
    browser = QTextBrowser()
    browser.setPlainText("First occurrence of quantum telemetry. Second occurrence of quantum telemetry.")
    find_bar = InPageFindBar(lang="en")
    find_bar.attach_browser(browser)

    find_bar.input_find.setText("quantum")
    find_bar.find_next()
    assert find_bar.count_label.text() == "2 matches"
    find_bar.find_prev()
    assert find_bar.count_label.text() == "2 matches"

    # Case sensitivity toggle
    find_bar.input_find.setText("Quantum")
    find_bar.chk_case.setChecked(True)
    find_bar._update_matches_count()
    assert "no matches" in find_bar.count_label.text().lower()
    find_bar.chk_case.setChecked(False)

    # 2. Citation Hover Tooltip Test
    doc = SourceDocument(cite_id=1, title="Quantum Telemetry 2026", content="Detailed empirical test logs.")
    builder = create_semantic_anchor_builder([doc])
    anchor_html = builder("1", page=None, context_claim="Detailed empirical test logs.")
    assert 'title="출처: [1] Quantum Telemetry 2026 | 발췌:' in anchor_html
    assert 'class="citation-anchor cite-verified"' in anchor_html

    # 3. BibTeX export test
    with tempfile.NamedTemporaryFile(suffix=".bib", delete=False) as tf:
        bib_path = Path(tf.name)

    doc2 = SourceDocument(
        cite_id=2,
        title="Neutral Atom Blueprint",
        path_or_uri="https://arxiv.org/abs/2601.12345",
        content="Blueprint for modular quantum architecture.",
    )
    export_bibtex_report(
        output_path=bib_path,
        markdown_text="Report citing [2].",
        documents=[doc, doc2],
    )
    bib_content = bib_path.read_text(encoding="utf-8")
    assert "@misc{source_2," in bib_content
    assert "title = {{Neutral Atom Blueprint}}" in bib_content
    assert "\\url{https://arxiv.org/abs/2601.12345}" in bib_content
    bib_path.unlink()

    print("   -> In-page find bar (Ctrl+F), hover tooltips, and BibTeX exporter verified successfully!")


def test_7_source_authority_and_drag_drop():
    print("[7/7] Testing Source Authority Hierarchy & Drag-and-Drop Event Handling...")

    # 1. Authority classification
    auth_gov = classify_source_authority("https://nist.gov/quantum", "NIST Quantum Telemetry")
    assert auth_gov == "PRIMARY/EMPIRICAL"

    auth_mil = classify_source_authority("https://darpa.mil/news", "DARPA Project Log")
    assert auth_mil == "PRIMARY/EMPIRICAL"

    auth_edu = classify_source_authority("https://harvard.edu/research", "Harvard Quantum Lab")
    assert auth_edu == "SECONDARY/SCHOLARLY"

    auth_arxiv = classify_source_authority("https://arxiv.org/abs/2602.001", "ArXiv Preprint")
    assert auth_arxiv == "SECONDARY/SCHOLARLY"

    auth_blog = classify_source_authority("https://medium.com/tech-review", "Opinion Piece")
    assert auth_blog == "TERTIARY/COMMENTARY"

    # 2. Drag and Drop on MainWindow
    config = ConfigManager()
    win = MainWindow(config_manager=config)

    # Test dragEnterEvent with valid files
    mime = QMimeData()
    mime.setUrls([QUrl.fromLocalFile("test.docx"), QUrl.fromLocalFile("test.pdf")])
    drag_event = QDragEnterEvent(
        QPointF(100.0, 100.0).toPoint(),
        Qt.DropAction.CopyAction,
        mime,
        Qt.MouseButton.LeftButton,
        Qt.KeyboardModifier.NoModifier,
    )
    win.dragEnterEvent(drag_event)
    assert drag_event.isAccepted(), "Valid documents must be accepted on dragEnter"

    # Test dropEvent with temporary document
    with tempfile.NamedTemporaryFile(suffix=".txt", delete=False) as tf:
        tf.write(b"Drag and drop text document payload.")
        drop_txt = Path(tf.name)

    drop_mime = QMimeData()
    drop_mime.setUrls([QUrl.fromLocalFile(str(drop_txt))])
    drop_event = QDropEvent(
        QPointF(100.0, 100.0),
        Qt.DropAction.CopyAction,
        drop_mime,
        Qt.MouseButton.LeftButton,
        Qt.KeyboardModifier.NoModifier,
    )
    initial_count = len(win.corpus.documents)
    win.dropEvent(drop_event)
    assert len(win.corpus.documents) == initial_count + 1, "Dropped file must be ingested into corpus"
    drop_txt.unlink()

    # Test shortcuts
    assert hasattr(win, "shortcut_find")
    assert hasattr(win, "shortcut_search_fts")
    assert hasattr(win, "search_dock")
    assert win.search_dock.isWidgetType()

    print("   -> Source authority hierarchy classification and MainWindow drag-and-drop verified successfully!")


if __name__ == "__main__":
    print("=== OMNIRESEARCH STUDIO PHASE 6 VERIFICATION SUITE ===")
    test_1_markdown_repair_and_continuation()
    test_2_audio_defenses_and_id3()
    test_3_context_caching()
    test_4_universal_document_ingestion()
    test_5_sqlite_fts5_and_search_drawer()
    test_6_find_bar_and_bibtex()
    test_7_source_authority_and_drag_drop()
    print("\nALL PHASE 6 TESTS COMPLETED SUCCESSFULLY!")
