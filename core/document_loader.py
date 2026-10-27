"""Document ingestion and indexing for PDF, TXT, and Markdown files."""

from __future__ import annotations

import datetime
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional

from pypdf import PdfReader


def load_pdf(file_path: str | Path) -> Dict[str, Any]:
    """
    Extract Markdown-native text from a PDF page-by-page using pymupdf4llm, inserting
    deterministic `<!--PAGE_START:{page_num}-->` page boundary markers, and automatically
    falling back to RapidOCR on scanned/image-only pages (<60 alphanumeric chars).
    """
    path = Path(file_path).resolve()
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    pages: List[str] = []
    has_ocr = False
    ocr_engine = None

    # Step 1: Attempt pymupdf4llm page-by-page markdown extraction
    chunks: List[Dict[str, Any]] = []
    try:
        import pymupdf4llm

        chunks = pymupdf4llm.to_markdown(str(path), page_chunks=True)
    except Exception:
        chunks = []

    # Step 2: Open with pymupdf for page counting and pixmap OCR fallback
    import pymupdf

    fitz_doc = None
    try:
        fitz_doc = pymupdf.open(str(path))
        total_pages = len(fitz_doc)
    except Exception:
        total_pages = len(chunks) if chunks else 0

    if total_pages == 0 and chunks:
        total_pages = len(chunks)

    # Process each page
    import hashlib
    import io
    from PIL import Image

    doc_hash = hashlib.sha256(str(path).encode("utf-8")).hexdigest()[:12]
    all_figures: List[Dict[str, Any]] = []
    from core.paths import get_cache_dir
    fig_dir = get_cache_dir() / "figures"
    fig_dir.mkdir(parents=True, exist_ok=True)

    for page_idx in range(total_pages):
        page_num = page_idx + 1
        page_md = ""

        if page_idx < len(chunks):
            page_md = str(chunks[page_idx].get("text", "")).strip()

        # Check alphanumeric density (scanned page threshold < 60 chars)
        alnum_count = sum(1 for ch in page_md if ch.isalnum())
        if alnum_count < 60 and fitz_doc is not None and page_idx < len(fitz_doc):
            try:
                if ocr_engine is None:
                    from rapidocr_onnxruntime import RapidOCR

                    ocr_engine = RapidOCR()

                page_obj = fitz_doc[page_idx]
                pix = page_obj.get_pixmap(dpi=200)
                ocr_res, _ = ocr_engine(pix.tobytes("png"))
                if ocr_res:
                    lines = [item[1] for item in ocr_res if len(item) > 1 and item[1]]
                    if lines:
                        page_md = "\n\n".join(lines)
                        has_ocr = True
            except Exception:
                pass

        # Extract embedded schematics & diagrams (width >= 150, height >= 150)
        page_figures: List[Dict[str, Any]] = []
        if fitz_doc is not None and page_idx < len(fitz_doc):
            try:
                page_obj = fitz_doc[page_idx]
                for img_idx, img_info in enumerate(page_obj.get_images(full=True)):
                    try:
                        xref = img_info[0]
                        base_img = fitz_doc.extract_image(xref)
                        w = base_img.get("width", 0)
                        h = base_img.get("height", 0)
                        if w >= 150 and h >= 150:
                            fig_name = f"{doc_hash}_p{page_num}_img{img_idx}.png"
                            fig_path = fig_dir / fig_name
                            img_bytes = base_img.get("image", b"")
                            pil_img = Image.open(io.BytesIO(img_bytes))
                            pil_img.save(fig_path, format="PNG")

                            fig_meta = {
                                "path": str(fig_path.resolve()),
                                "file_path": str(fig_path.resolve()),
                                "filename": fig_name,
                                "page_num": page_num,
                                "index": img_idx,
                                "width": w,
                                "height": h,
                            }
                            all_figures.append(fig_meta)
                            page_figures.append(fig_meta)
                    except Exception:
                        pass
            except Exception:
                pass

        figure_anchors = ""
        if page_figures:
            figure_anchors = "\n\n" + "\n".join(
                f"![Figure {f['page_num']}-{f['index']}](file:///{Path(f['path']).as_posix()})"
                for f in page_figures
            )

        # Prepend deterministic page boundary marker
        page_text = f"<!--PAGE_START:{page_num}-->\n\n{page_md}{figure_anchors}"
        pages.append(page_text)

    if fitz_doc is not None:
        try:
            fitz_doc.close()
        except Exception:
            pass

    # Fallback to pypdf if both pymupdf4llm and fitz failed to yield pages
    if not pages:
        try:
            reader = PdfReader(str(path))
            for page_idx, page_obj in enumerate(reader.pages):
                page_num = page_idx + 1
                txt = page_obj.extract_text() or ""
                pages.append(f"<!--PAGE_START:{page_num}-->\n\n{txt.strip()}")
            total_pages = len(pages)
        except Exception:
            total_pages = 0

    full_content = "\n\n".join(pages).strip()
    return {
        "content": full_content,
        "page_count": max(1, total_pages),
        "has_ocr": has_ocr,
        "pages": pages,
        "figures": all_figures,
    }


def load_docx(file_path: str | Path) -> Dict[str, Any]:
    """
    Ingest a .docx Word document, extracting text paragraphs and converting tables
    into Markdown pipe tables with deterministic <!--PAGE_START:1--> anchor.
    """
    path = Path(file_path).resolve()
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    import docx

    doc = docx.Document(str(path))
    blocks: List[str] = []

    for child in doc.element.body:
        if child.tag.endswith("p"):
            p = docx.text.paragraph.Paragraph(child, doc)
            text = p.text.strip()
            if text:
                blocks.append(text)
        elif child.tag.endswith("tbl"):
            table = docx.table.Table(child, doc)
            table_lines: List[str] = []
            for i, row in enumerate(table.rows):
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                table_lines.append("| " + " | ".join(cells) + " |")
                if i == 0:
                    table_lines.append("| " + " | ".join(["---"] * len(cells)) + " |")
            if table_lines:
                blocks.append("\n".join(table_lines))

    raw_text = "\n\n".join(blocks).strip()
    full_content = f"<!--PAGE_START:1-->\n\n{raw_text}"
    return {
        "content": full_content,
        "page_count": 1,
        "has_ocr": False,
        "pages": [full_content],
    }


AUDIO_EXTENSIONS = {".mp3", ".wav", ".m4a", ".aac", ".ogg", ".flac"}
DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown"}
SUPPORTED_EXTENSIONS = DOCUMENT_EXTENSIONS | AUDIO_EXTENSIONS


@dataclass
class SourceDocument:
    """Represents an indexed source document or discovered web perspective."""

    cite_id: int
    title: str
    source_type: str = "local_txt"  # "local_pdf", "local_txt", "local_md", "web_sweep"
    path_or_uri: str = ""
    content: str = ""
    word_count: int = 0
    page_count: int = 1
    enabled: bool = True
    metadata: Dict[str, object] = field(default_factory=dict)
    created_at: str = field(
        default_factory=lambda: datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    )

    def __post_init__(self) -> None:
        if not self.word_count and self.content:
            self.word_count = len(self.content.split())

    @property
    def citation_tag(self) -> str:
        """Formatted bracket citation marker."""
        return f"[{self.cite_id}]"

    @property
    def estimated_tokens(self) -> int:
        """Rough token estimate (~4 characters per token for Latin text, floor of 1.3 tokens per word)."""
        return max(int(len(self.content) / 4), int(self.word_count * 1.3), 1 if self.content else 0)

    def to_dict(self, include_content: bool = True) -> Dict[str, object]:
        """Serialize this document (content optional) for project persistence."""
        data: Dict[str, object] = {
            "cite_id": self.cite_id,
            "title": self.title,
            "source_type": self.source_type,
            "path_or_uri": self.path_or_uri,
            "word_count": self.word_count,
            "page_count": self.page_count,
            "enabled": self.enabled,
            "metadata": dict(self.metadata),
            "created_at": self.created_at,
        }
        if include_content:
            data["content"] = self.content
        return data

    @classmethod
    def from_dict(cls, data: Dict[str, object], content: Optional[str] = None) -> "SourceDocument":
        """Rebuild a document from `to_dict` output; `content` overrides the dict value when provided."""
        text = content if content is not None else str(data.get("content", ""))
        metadata = data.get("metadata")
        return cls(
            cite_id=int(data.get("cite_id", 0)),  # type: ignore[arg-type]
            title=str(data.get("title", "")),
            source_type=str(data.get("source_type", "local_txt")),
            path_or_uri=str(data.get("path_or_uri", "")),
            content=text,
            word_count=int(data.get("word_count", len(text.split()))),  # type: ignore[arg-type]
            page_count=int(data.get("page_count", 1)),  # type: ignore[arg-type]
            enabled=bool(data.get("enabled", True)),
            metadata=dict(metadata) if isinstance(metadata, dict) else {},
            created_at=str(data.get("created_at", "")) or datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        )

    @property
    def summary_snippet(self) -> str:
        """Short preview of the document text."""
        cleaned = re.sub(r"\s+", " ", self.content).strip()
        if len(cleaned) <= 220:
            return cleaned
        return cleaned[:217] + "..."


class DocumentCorpus:
    """Manages local documents and OmniSweep web research documents in a unified corpus."""

    def __init__(self) -> None:
        self._documents: List[SourceDocument] = []
        self._next_id: int = 1

    @property
    def documents(self) -> List[SourceDocument]:
        """Return all documents in the corpus."""
        return list(self._documents)

    @property
    def enabled_documents(self) -> List[SourceDocument]:
        """Return only documents currently enabled for synthesis."""
        return [doc for doc in self._documents if doc.enabled]

    def get_by_cite_id(self, cite_id: int) -> Optional[SourceDocument]:
        """Retrieve a document by its integer citation ID."""
        for doc in self._documents:
            if doc.cite_id == cite_id:
                return doc
        return None

    def get_document(self, cite_id: int) -> Optional[SourceDocument]:
        """Convenience alias for get_by_cite_id."""
        return self.get_by_cite_id(cite_id)

    def get_all_documents(self) -> List[SourceDocument]:
        """Convenience alias for self.documents."""
        return list(self._documents)

    def add_document(self, doc: SourceDocument) -> SourceDocument:
        """Add an already constructed SourceDocument directly into the corpus."""
        self._documents.append(doc)
        self._next_id = max(self._next_id, doc.cite_id + 1)
        return doc

    def remove_document(self, cite_id: int) -> bool:
        """Remove a document by citation ID."""
        for idx, doc in enumerate(self._documents):
            if doc.cite_id == cite_id:
                del self._documents[idx]
                return True
        return False

    def clear(self) -> None:
        """Clear all documents and reset citation numbering."""
        self._documents.clear()
        self._next_id = 1

    def restore_documents(self, documents: List[SourceDocument]) -> None:
        """Replace the corpus with `documents`, preserving their original citation IDs."""
        self._documents = list(documents)
        self._next_id = max((d.cite_id for d in self._documents), default=0) + 1

    def set_enabled(self, cite_id: int, enabled: bool) -> None:
        """Toggle whether a document is included in synthesis."""
        doc = self.get_by_cite_id(cite_id)
        if doc is not None:
            doc.enabled = enabled

    def ingest_file(self, file_path: str | Path, use_sandbox: bool = True) -> SourceDocument:
        """Load a .pdf, .docx, .txt, or .md file from disk and index it into the corpus using process sandboxing."""
        path = Path(file_path).resolve()
        if not path.exists() or not path.is_file():
            raise FileNotFoundError(f"File not found: {path}")

        ext = path.suffix.lower()
        if ext in AUDIO_EXTENSIONS:
            return self.ingest_audio_file(path)
        if ext not in SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"Unsupported file extension '{ext}'. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
            )

        has_ocr = False
        figures: List[Dict[str, Any]] = []

        if use_sandbox:
            from core.sandbox_loader import SandboxedDocumentLoader
            loader = SandboxedDocumentLoader.get_instance()
            res = loader.extract_document(path)
            if not res.get("success", False):
                title_str = res.get("title", f"[EXTRACTION_FAILED: Corrupted Native Binary/Bitmap] {path.name}")
                content = res.get("content", f"[Extraction failed: {res.get('error', 'Unknown error')}]")
                doc = SourceDocument(
                    cite_id=self._next_id,
                    title=title_str,
                    source_type="local_error",
                    path_or_uri=str(path),
                    content=content,
                    word_count=len(content.split()),
                    page_count=1,
                    enabled=False,
                    metadata={
                        "file_size_bytes": path.stat().st_size,
                        "extension": ext,
                        "has_ocr": False,
                        "figures": [],
                        "crashed": res.get("crashed", False),
                        "error": res.get("error", ""),
                    },
                )
                self._next_id += 1
                self._documents.append(doc)
                return doc

            text = res.get("content", "")
            page_count = res.get("page_count", 1)
            has_ocr = res.get("has_ocr", False)
            figures = res.get("figures", [])
            source_type = res.get("source_type", "local_txt")
        else:
            if ext == ".pdf":
                text, page_count, has_ocr, figures = self._extract_pdf(path)
                source_type = "local_pdf"
            elif ext == ".docx":
                text, page_count = self._extract_docx(path)
                source_type = "local_docx"
            elif ext in {".md", ".markdown"}:
                text = self._extract_text(path)
                page_count = 1
                source_type = "local_md"
            else:
                text = self._extract_text(path)
                page_count = 1
                source_type = "local_txt"

        normalized = self._normalize_whitespace(text)
        if not normalized:
            normalized = f"[Document '{path.name}' contained no extractable text.]"

        words = len(normalized.split())
        title_str = f"[OCR Scanned] {path.name}" if has_ocr else path.name
        doc = SourceDocument(
            cite_id=self._next_id,
            title=title_str,
            source_type=source_type,
            path_or_uri=str(path),
            content=normalized,
            word_count=words,
            page_count=page_count,
            enabled=True,
            metadata={"file_size_bytes": path.stat().st_size, "extension": ext, "has_ocr": has_ocr, "figures": figures},
        )
        self._next_id += 1
        self._documents.append(doc)
        return doc

    def ingest_audio_file(self, file_path: str | Path, progress_callback: Optional[Any] = None) -> SourceDocument:
        """Transcribe local audio media using faster-whisper and index into corpus with timestamp grounding."""
        from core.media_loader import MediaLoader
        path = Path(file_path).resolve()
        res = MediaLoader.transcribe_local_audio(path, progress_callback=progress_callback)
        content = res["content"]
        words = len(content.split())
        doc = SourceDocument(
            cite_id=self._next_id,
            title=res["title"],
            source_type="local_audio",
            path_or_uri=str(path),
            content=content,
            word_count=words,
            page_count=1,
            enabled=True,
            metadata={
                "file_size_bytes": path.stat().st_size if path.exists() else 0,
                "extension": path.suffix.lower(),
                "duration_seconds": res.get("duration_seconds", 0.0),
                **(res.get("metadata", {})),
            },
        )
        self._next_id += 1
        self._documents.append(doc)
        return doc

    def ingest_youtube_url(self, url: str) -> SourceDocument:
        """Ingest YouTube video transcripts with timestamp anchors into the corpus."""
        from core.media_loader import MediaLoader
        res = MediaLoader.ingest_youtube_url(url)
        content = res["content"]
        words = len(content.split())
        doc = SourceDocument(
            cite_id=self._next_id,
            title=res["title"],
            source_type="web_youtube",
            path_or_uri=res["url"],
            content=content,
            word_count=words,
            page_count=1,
            enabled=True,
            metadata=res.get("metadata", {}),
        )
        self._next_id += 1
        self._documents.append(doc)
        return doc

    def ingest_academic(self, identifier_or_url: str, download_pdf: bool = True) -> SourceDocument:
        """Resolve ArXiv or DOI into indexed document or downloaded vector PDF."""
        from core.academic_resolver import AcademicResolver
        res = AcademicResolver.resolve(identifier_or_url, download_pdf=download_pdf)
        pdf_path = res.get("pdf_path")
        if pdf_path and Path(pdf_path).exists():
            doc = self.ingest_file(pdf_path)
            doc.title = res["title"]
            doc.source_type = "academic_arxiv" if "arxiv" in str(doc.path_or_uri).lower() else "academic_doi"
            hdr = res.get("header_markdown", "")
            if hdr:
                doc.content = f"{hdr}\n\n{doc.content}"
                doc.word_count = len(doc.content.split())
            doc.metadata.update(res.get("metadata", {}))
            return doc
        else:
            content = res.get("header_markdown", "") or res.get("abstract", "")
            doc = SourceDocument(
                cite_id=self._next_id,
                title=res["title"],
                source_type=res.get("source_type", "academic_doi"),
                path_or_uri=res.get("doi") or res.get("arxiv_id") or identifier_or_url,
                content=content,
                word_count=len(content.split()),
                page_count=1,
                enabled=True,
                metadata=res.get("metadata", {}),
            )
            self._next_id += 1
            self._documents.append(doc)
            return doc

    def add_web_document(
        self,
        title: str,
        uri: str,
        content: str,
        perspective_name: str = "",
        grounding_links: Optional[List[Dict[str, str]]] = None,
        origin_tag: str = "",
        extra_metadata: Optional[Dict[str, object]] = None,
    ) -> SourceDocument:
        """Add a discovered web perspective or grounded web source into the corpus with linguistic origin tagging."""
        from core.storm_discovery import (
            detect_linguistic_origin_tag,
            format_multilingual_source_tag,
        )

        normalized = self._normalize_whitespace(content)
        resolved_origin = origin_tag or detect_linguistic_origin_tag(
            f"{title} {normalized}",
            uri=uri,
            perspective_name=perspective_name,
        )
        already_tagged = title.lstrip().startswith("[WEB ")
        tagged_title = (
            format_multilingual_source_tag(perspective_name, resolved_origin, title)
            if perspective_name and not already_tagged
            else title
        )
        words = len(normalized.split())
        doc = SourceDocument(
            cite_id=self._next_id,
            title=tagged_title,
            source_type="web_sweep",
            path_or_uri=uri,
            content=normalized,
            word_count=words,
            page_count=1,
            enabled=True,
            metadata={
                "perspective": perspective_name,
                "origin_tag": resolved_origin,
                "grounding_links": grounding_links or [],
                **(extra_metadata or {}),
            },
        )
        self._next_id += 1
        self._documents.append(doc)
        return doc

    @staticmethod
    def _extract_pdf(path: Path) -> tuple[str, int, bool, List[Dict[str, Any]]]:
        """Extract text from all pages of a PDF file using load_pdf."""
        res = load_pdf(path)
        return res["content"], res["page_count"], res["has_ocr"], res.get("figures", [])

    @staticmethod
    def _extract_docx(path: Path) -> tuple[str, int]:
        """Extract text and markdown tables from a docx file using load_docx."""
        res = load_docx(path)
        return res["content"], res["page_count"]

    @staticmethod
    def _extract_text(path: Path) -> str:
        """Read plain text or markdown with chardet encoding detection (UTF-8, CP949, EUC-KR, CP1252, etc.)."""
        import chardet

        raw_bytes = path.read_bytes()
        detected = chardet.detect(raw_bytes)
        det_enc = detected.get("encoding") if detected else None

        candidates = [det_enc, "utf-8", "utf-8-sig", "cp949", "euc-kr", "cp1252", "latin-1"]
        extracted = None
        for enc in candidates:
            if not enc:
                continue
            try:
                extracted = raw_bytes.decode(enc)
                break
            except (UnicodeDecodeError, LookupError):
                continue

        if extracted is None:
            extracted = raw_bytes.decode("utf-8", errors="replace")

        if not extracted.startswith("<!--PAGE_START:"):
            extracted = f"<!--PAGE_START:1-->\n\n{extracted}"

        return extracted

    @staticmethod
    def _normalize_whitespace(text: str) -> str:
        """Collapse excessive blank lines while preserving paragraph structure."""
        text = text.replace("\r\n", "\n").replace("\r", "\n")
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()
